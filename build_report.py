"""PPT editable, grafik evaluasi aktual, dan outline Markdown dari artefak project."""
import json
from pathlib import Path

import joblib
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from sklearn.calibration import calibration_curve
from sklearn.metrics import ConfusionMatrixDisplay, RocCurveDisplay

from prediction import example_customers, input_schema, predict_customers
from training_data import load_training_split

BASE = Path(__file__).resolve().parent
DOCS = BASE / "docs"
FIG = DOCS / "figures"
APP_URL = "https://machine-learning-project-u8yovshrqrc5tm5s2y2qto.streamlit.app/"
PPT_NAME = "ML2026_B2_Kelompok12_PrediksiCustomerChurn.pptx"
BG, CARD, TEXT, MUTED, ACCENT = "0B0F19", "141B2E", "E9EDF5", "AAB4C8", "7C3AED"


def rgb(value):
    return RGBColor.from_string(value)


def evaluation_figures():
    FIG.mkdir(parents=True, exist_ok=True)
    comparison = pd.read_csv(BASE / "hasil_perbandingan.csv")
    deployment = json.loads((BASE / "deployment_info.json").read_text())
    pipe, cal = joblib.load(BASE / "model_churn.pkl"), joblib.load(BASE / "calibrator.pkl")
    _, X, _, y = load_training_split()
    probability, predictions = predict_customers(X, cal, input_schema(pipe), deployment["threshold"])
    plt.rcParams.update({"font.size": 11, "figure.facecolor": "white"})

    def save(fig, name):
        fig.tight_layout()
        fig.savefig(FIG / name, dpi=170, bbox_inches="tight")
        plt.close(fig)

    fig, ax = plt.subplots(figsize=(9, 4))
    comparison.set_index("Model")[["Accuracy", "Precision", "Recall", "F1"]].plot.bar(
        ax=ax, color=["#7c3aed", "#2563eb", "#06b6d4", "#10b981"], rot=0)
    ax.set(ylim=(0, 1), ylabel="Skor test", title="Tiga model inti — threshold 0,50")
    ax.legend(ncol=4, fontsize=9)
    save(fig, "model_comparison.png")
    fig, ax = plt.subplots(figsize=(9, 4))
    comparison.set_index("Model")[["Train_F1", "CV_F1", "F1"]].plot.bar(
        ax=ax, color=["#7c3aed", "#2563eb", "#10b981"], rot=0)
    ax.set(ylim=(0, 1), ylabel="F1 kelas churn", title="Train / validation CV / test")
    ax.legend(["Train penuh", "Validation CV", "Test"], ncol=3)
    save(fig, "overfitting.png")
    fig, ax = plt.subplots(figsize=(6, 4.5))
    ConfusionMatrixDisplay.from_predictions(y, predictions, display_labels=["Tidak churn", "Churn"],
                                           cmap="Purples", colorbar=False, ax=ax)
    ax.set_title("Model aplikasi — test 1.407 pelanggan")
    save(fig, "confusion_matrix_app.png")
    fig, axes = plt.subplots(1, 2, figsize=(10, 4))
    RocCurveDisplay.from_predictions(y, probability, ax=axes[0], name="Aplikasi", curve_kwargs={"color": "#7c3aed"})
    axes[0].plot([0, 1], [0, 1], "--", color="#94a3b8")
    axes[0].set_title("ROC — holdout test")
    for p, name, color in [(pipe.predict_proba(X)[:, 1], "Model inti", "#ef4444"),
                           (probability, "Terkalibrasi", "#2563eb")]:
        frac, mean = calibration_curve(y, p, n_bins=8)
        axes[1].plot(mean, frac, "o-", label=name, color=color)
    axes[1].plot([0, 1], [0, 1], "--", color="#94a3b8")
    axes[1].set(xlabel="Rata-rata probabilitas", ylabel="Proporsi churn", title="Reliability curve — test")
    axes[1].legend()
    save(fig, "roc_calibration.png")
    selection = pd.read_csv(BASE / "threshold_selection.csv")
    fig, ax = plt.subplots(figsize=(8, 3.6))
    ax.plot(selection.Threshold, selection.OOF_F1, color="#7c3aed", linewidth=2)
    ax.axvline(deployment["threshold"], linestyle="--", color="#10b981", label="Threshold terpilih 0,32")
    ax.set(xlabel="Threshold", ylabel="OOF F1 training (skor tuning)", title="Pemilihan threshold — train only")
    ax.legend()
    save(fig, "threshold_train.png")
    coefficients = pd.DataFrame({"Fitur": pipe.named_steps["pre"].get_feature_names_out(),
                                 "Koefisien": pipe.named_steps["clf"].coef_[0]})
    top = coefficients.reindex(coefficients.Koefisien.abs().sort_values(ascending=False).index).head(10)
    top = top.sort_values("Koefisien")
    fig, ax = plt.subplots(figsize=(9, 4.5))
    ax.barh(top.Fitur.str.replace(r"^(num|cat)__", "", regex=True), top.Koefisien,
            color=["#ef4444" if c > 0 else "#10b981" for c in top.Koefisien])
    ax.set(xlabel="Koefisien log-odds (model inti)", title="Hubungan fitur bersyarat, bukan sebab-akibat")
    save(fig, "coefficients.png")
    example = example_customers()
    p, label = predict_customers(example, cal, input_schema(pipe), deployment["threshold"])
    example["Prob_Churn"] = (p * 100).round(2)
    example["Prediksi"] = np.where(label == 1, "CHURN", "TETAP")
    example.to_csv(DOCS / "contoh_demo.csv", index=False)
    example_customers().to_csv(DOCS / "template_pelanggan.csv", index=False)
    return comparison, deployment, example


def build_report():
    comparison, deploy, example = evaluation_figures()
    summary = json.loads((DOCS / "eda" / "summary.json").read_text(encoding="utf-8"))
    refs = json.loads((DOCS / "referensi.json").read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = "Prediksi Customer Churn — Telco"
    prs.core_properties.author = "Kelompok 12 — Kelas B2"
    outlines = []

    def text(slide, value, x, y, w, h, size=20, color=TEXT, bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.02)
        tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(str(value).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name, p.font.size, p.font.bold, p.font.color.rgb = "Aptos", Pt(size), bold, rgb(color)
            p.space_after = Pt(10)
        return box

    def slide(title, subtitle="", notes=""):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(BG)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
        bar.fill.solid(); bar.fill.fore_color.rgb = rgb(ACCENT); bar.line.fill.background()
        text(s, title, 0.55, 0.42, 12.2, 0.65, size=29, bold=True)
        if subtitle:
            text(s, subtitle, 0.58, 1.13, 12.1, 0.50, size=14, color=MUTED)
        text(s, f"ML 2026  •  B2  •  Kelompok 12                                      {len(prs.slides):02d}",
             0.6, 7.09, 12.1, 0.28, size=11, color=MUTED)
        s.notes_slide.notes_text_frame.text = notes or title
        outlines.append({"slide": len(prs.slides), "title": title, "subtitle": subtitle, "notes": notes})
        return s

    def bullets(s, lines, x=0.65, y=1.85, w=12.0, h=4.8, size=22):
        text(s, "\n".join("• " + line for line in lines), x, y, w, h, size=size)

    def picture(s, path, x, y, w, h):
        with Image.open(path) as image:
            iw, ih = image.size
        scale = min(w / iw, h / ih)
        pw, ph = iw * scale, ih * scale
        s.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                             width=Inches(pw), height=Inches(ph))

    def table(s, headers, rows, x=0.6, y=1.85, w=12.1, h=3.8, size=17, widths=None):
        shape = s.shapes.add_table(len(rows)+1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
        t = shape.table
        if widths:
            for col, fraction in zip(t.columns, widths):
                col.width = Inches(w * fraction)
        for i, values in enumerate([headers] + rows):
            for j, value in enumerate(values):
                cell = t.cell(i, j)
                cell.text = str(value)
                cell.fill.solid(); cell.fill.fore_color.rgb = rgb(ACCENT if i == 0 else CARD)
                cell.margin_left = cell.margin_right = Inches(0.10)
                for p in cell.text_frame.paragraphs:
                    p.font.name = "Aptos"; p.font.size = Pt(size); p.font.color.rgb = rgb(TEXT)
                    p.font.bold = (i == 0)
        return shape

    s = slide("Prediksi Customer Churn — Telco", "Perbandingan tiga model dan aplikasi Streamlit terkalibrasi",
              "Perkenalkan seluruh anggota. Judul project, Kelompok 12 Kelas B2. Buka aplikasi live pada bagian demo.")
    text(s, "KELOMPOK 12  /  KELAS B2", 0.7, 2.0, 11, 0.5, size=21, color="A78BFA", bold=True)
    text(s, "Muhamad Akhdan Ramadhan  —  J0404241102\nThevan Erlangga  —  J0404241073\n"
         "Fachri Abyasa Tarid  —  J0404241136", 0.7, 2.85, 11.8, 2.0, size=25)
    text(s, "Mata Kuliah Pembelajaran Mesin • Project Akhir 2026", 0.7, 5.8, 11.8, 0.5, size=18, color=MUTED)

    s = slide("Latar belakang & rumusan masalah", "Churn = pelanggan berhenti berlangganan")
    bullets(s, ["Pelanggan memiliki profil kontrak, layanan, dan pembayaran yang berbeda.",
                "Identifikasi risiko membantu menentukan prioritas tindak lanjut retensi [1, 4].",
                "Masalah: bagaimana memprediksi churn dari data pelanggan secara terukur?",
                "Tantangan: kelas churn minoritas dan biaya false alarm / churn terlewat berbeda.",
                "Prediksi risiko belum membuktikan efektivitas tindakan retensi."])
    s = slide("Tujuan project", "Klasifikasi supervised tabular, bukan forecasting")
    bullets(s, ["Memahami dataset melalui EDA dan preprocessing yang sesuai.",
                "Membandingkan Logistic Regression, Decision Tree, dan Random Forest.",
                "Memilih model inti melalui validation CV-F1 pada data training.",
                "Menghasilkan probabilitas terkalibrasi dan keputusan dengan threshold train-only.",
                "Menyediakan aplikasi single/batch CSV dengan validasi dan dashboard transparan."])

    for group, title in [(refs[:3], "Studi literatur — metode & evaluasi"), (refs[3:], "Studi literatur — early warning & pengembangan")]:
        s = slide(title, "Lima jurnal terverifikasi • ringkasan hasil penulis, bukan hasil project")
        rows = []
        for r in group:
            dataset = {1: "4 dataset; termasuk Telco 7.043", 2: "Orange 3.333", 3: "GitHub 5.000 / BigML 3.333",
                       4: "Operator 900.000 mentah", 5: "Telco; 7.033 menurut artikel"}[r["no"]]
            method = {1: "8 model; transformasi + feature selection + grid search", 2: "Kernel SVM + SFS/SBS + SMOTE-ENN",
                      3: "Clustering + classifier ensemble", 4: "RF-AdaBoost + SMOTE", 5: "BiLSTM-CNN"}[r["no"]]
            result = {1: "WOE+LR: AUC 0,796 / F1 0,79 (Dataset-1)", 2: "Accuracy 0,9888 / F1 0,9901",
                      3: "Accuracy 94,7% / 92,43%", 4: "P=0,99 / R=0,93 / F1=0,96 (churn)",
                      5: "Accuracy 81% / F-score 65%"}[r["no"]]
            rows.append([f"[{r['no']}] {r['short']}", dataset, method, result])
        table(s, ["Referensi", "Dataset", "Metode", "Hasil utama"], rows, h=3.6, size=16,
              widths=[.22, .23, .28, .27])
        text(s, "Project: 3 model tabular + pipeline train-only + sigmoid calibration + Streamlit. "
             "Skor jurnal tidak dibandingkan langsung karena dataset dan protokol berbeda.", 0.65, 5.8, 12, 0.9, size=18, color=MUTED)
    s = slide("Dataset & target", "Sumber: IBM Telco Customer Churn di Kaggle — blastchar")
    table(s, ["Komponen", "Jumlah / karakteristik"], [
        ["Data mentah", "7.043 pelanggan • 21 kolom (ID + 19 kandidat fitur + target)"],
        ["Cohort pemodelan", "7.032 pelanggan • 18 input fitur + target"],
        ["Target", "Churn Yes=1 / No=0 • 1.869 churn dan 5.163 tidak churn"],
        ["Numerik", "tenure (bulan), MonthlyCharges ($); SeniorCitizen indikator 0/1"],
        ["Kategorikal", "15 fitur: kontrak, layanan, pembayaran, demografi"],
    ], h=3.9, size=19, widths=[.27,.73])
    text(s, "Sumber dan kamus seluruh 21 variabel: docs/eda/EDA.md", .65, 6.05, 12, .5, size=17, color=MUTED)
    s = slide("Pemeriksaan kualitas data", "Missing, duplikat, outlier, dan keputusan cleaning")
    bullets(s, ["11 TotalCharges kosong → numeric coercion → missing; seluruhnya tenure 0 / tidak churn.",
                "Tidak ada customerID ganda atau baris mentah identik.",
                "27 profil duplikat setelah ID/TotalCharges dihapus (termasuk target); tidak dibuang otomatis.",
                "IQR 1,5× pada tenure, MonthlyCharges, TotalCharges: 0 outlier; tidak dilakukan clipping.",
                "18 profil test cocok dengan profil train; menjadi keterbatasan generalisasi profil baru."], size=21)

    for title, file, note in [
        ("EDA — distribusi kelas", "01_target_distribution.png", summary["figures"][0]["interpretation"]),
        ("EDA — distribusi numerik", "02_numeric_histograms_train.png", summary["figures"][1]["interpretation"]),
        ("EDA — profil churn vs tidak churn", "03_numeric_boxplots_train.png", summary["figures"][2]["interpretation"]),
        ("EDA — kontrak, internet, pembayaran", "05_categorical_churn_train.png", summary["figures"][4]["interpretation"]),
    ]:
        s = slide(title, "Hubungan fitur–target dianalisis pada training set saja", notes=note)
        picture(s, DOCS / "eda" / file, .6, 1.8, 12.1, 3.6)
        text(s, note, .65, 5.60, 12, 1.25, size=16, color=MUTED)
    s = slide("EDA — korelasi & pilihan fitur", "Korelasi bukan identitas dan bukan sebab-akibat")
    picture(s, DOCS / "eda" / "04_correlation_train.png", .6, 1.8, 6.4, 4.8)
    bullets(s, ["TotalCharges berkorelasi kuat dengan tenure.", "TotalCharges dihapus agar input aplikasi sederhana.",
                "Tidak selalu MonthlyCharges × tenure.", "Belum ada ablation untuk membuktikan keuntungan penghapusan fitur."],
            x=7.25, y=2, w=5.4, h=4.8, size=20)
    s = slide("Preprocessing & pencegahan leakage", "Fit preprocessing hanya pada training fold")
    bullets(s, ["customerID dihapus; TotalCharges tidak digunakan; 11 catatan tidak lengkap dieksklusi.",
                "StandardScaler: SeniorCitizen, tenure, MonthlyCharges; OneHotEncoder: 15 kategori.",
                "Scaling membantu LR; pohon tetap memakai pipeline konsisten walau tidak memerlukannya.",
                "Class weight dituning untuk imbalance; tanpa SMOTE pada eksperimen project.",
                "ColumnTransformer + classifier dalam Pipeline; test tidak dipakai untuk fit/tuning."], size=21)
    s = slide("Tiga algoritma & alasan pemilihan", "Membandingkan model linear, pohon tunggal, dan ensemble")
    table(s, ["Model", "Alasan", "Risiko / kontrol"], [
        ["Logistic Regression", "Ringan; baseline linear log-odds; koefisien mudah dibahas", "C dan class_weight dituning"],
        ["Decision Tree", "Aturan non-linear dan interaksi fitur", "max_depth / min_samples_split / class_weight"],
        ["Random Forest", "Bagging banyak pohon; pembanding kompleksitas", "Kedalaman dan split dituning; periksa overfitting"],
    ], h=3.7, size=20, widths=[.24,.42,.34])
    text(s, "Dasar literatur: [1] membandingkan ketiganya; [3]–[4] membahas ensemble churn.", .65, 5.95, 12, .6, size=18, color=MUTED)
    s = slide("Training, validation & parameter terbaik", "80:20 stratified • random_state=42 • CV 5-fold • GridSearch scoring F1")
    table(s, ["Model", "Parameter terbaik", "CV-F1"], [
        ["Logistic Regression", "C=10; class_weight=balanced; max_iter=1.000", "0,6304"],
        ["Decision Tree", "max_depth=5; min_samples_split=2; balanced", "0,6147"],
        ["Random Forest", "200 pohon; depth=12; min_samples_split=5; balanced_subsample", "0,6260"],
    ], h=3.5, size=19, widths=[.24,.61,.15])
    text(s, "Train 5.625; test 1.407. Grid: LR 6, DT 12, RF 8 konfigurasi × 5 fold = 130 fit CV, "
         "ditambah refit pemenang setiap model. Test tidak masuk fold training.", .65, 5.65, 12, 1, size=18, color=MUTED)
    s = slide("Evaluasi & perbandingan model inti", "Label positif: churn • threshold 0,50 • holdout yang sama")
    rows = [[r.Model, f"{r.CV_F1:.4f}", f"{r.Accuracy:.4f}", f"{r.Precision:.4f}",
             f"{r.Recall:.4f}", f"{r.F1:.4f}", f"{r.ROC_AUC:.4f}"] for r in comparison.itertuples()]
    table(s, ["Model", "CV-F1", "Accuracy", "Precision", "Recall", "F1", "AUC"], rows, h=2.6, size=18,
          widths=[.25,.125,.125,.125,.125,.125,.125])
    text(s, "LR dipilih dengan CV-F1 tertinggi pada train. DT memiliki recall/F1 test tertinggi; RF accuracy "
         "tertinggi. DT CV-F1 terendah; LR F1 test terendah di antara model inti. "
         "Selisih CV-F1 LR–RF hanya 0,0044, bukan bukti unggul mutlak.", .65, 4.9, 12, 1.5, size=20)
    s = slide("Analisis model & overfitting", "Gunakan train–validation–test, bukan accuracy saja")
    picture(s, FIG / "overfitting.png", .6, 1.8, 7.1, 4.4)
    bullets(s, ["RF: F1 train 0,8217 vs test 0,6096 → indikasi overfitting.",
                "LR: train 0,6338 vs test 0,6043; gap lebih kecil.",
                "Baseline selalu tidak churn: accuracy 73,42%, recall churn 0.",
                "Kompleksitas, regularisasi, class_weight dan interaksi dapat memengaruhi hasil."],
            x=8, y=1.9, w=4.7, h=4.9, size=18)
    s = slide("Kalibrasi & pemilihan threshold", "Model aplikasi: LR + sigmoid calibration 5-fold, train-only")
    picture(s, FIG / "threshold_train.png", .6, 1.8, 7.1, 3.7)
    bullets(s, ["Grid threshold 0,10–0,60; langkah 0,01.", "F1 maksimum prediksi OOF train → threshold 0,32.",
                "Kalibrasi di dalam setiap outer training fold.", "OOF F1 0,6338 adalah skor tuning, bukan estimasi independen.",
                "CHURN jika probabilitas ≥32%; TETAP jika lebih rendah."], x=8, y=1.9, w=4.7, h=4.9, size=18)
    text(s, "Hyperparameter inti telah dipilih pada train-CV. Test hanya dipakai untuk evaluasi akhir.", .65, 6.2, 12, .5, size=17, color=MUTED)
    s = slide("Evaluasi konfigurasi aplikasi", "Probabilitas dan keputusan berasal dari model yang sama")
    metrics = deploy["metrics"]
    table(s, ["Konfigurasi", "Accuracy", "Precision", "Recall", "F1", "AUC", "Brier"], [
        ["Threshold 0,32", *[f"{metrics[k]:.4f}" for k in ["Accuracy","Precision","Recall","F1","ROC_AUC","Brier"]]],
        ["Threshold 0,50", *[f"{deploy['metrics_at_0_5'][k]:.4f}" for k in ["Accuracy","Precision","Recall","F1","ROC_AUC","Brier"]]],
    ], h=2.5, size=19, widths=[.25,.125,.125,.125,.125,.125,.125])
    text(s, "Threshold 0,32 mendeteksi 275 churn vs 211 pada 0,50; false alarm naik 115 → 243. "
         "Trade-off precision–recall; ROC-AUC/Brier tidak berubah jika hanya threshold diganti.", .65, 4.8, 12, 1.2, size=21)
    s = slide("Confusion matrix & kesalahan prediksi", "Test: 1.033 tidak churn dan 374 churn")
    picture(s, FIG / "confusion_matrix_app.png", .6, 1.7, 6.4, 4.9)
    bullets(s, ["TP 275: churn berhasil terdeteksi.", "FN 99: churn terlewat.", "FP 243: pelanggan tidak churn ditandai churn.",
                "TN 790: tidak churn diprediksi benar.", "Precision 53,09%; recall 73,53%. Biaya retensi nyata belum diukur."],
            x=7.3, y=2, w=5.4, h=4.8, size=21)
    s = slide("ROC & kualitas probabilitas", "Kurva dihitung dari model tersimpan pada holdout test")
    picture(s, FIG / "roc_calibration.png", .6, 1.8, 12.1, 4.2)
    text(s, "AUC aplikasi 0,8324. Brier turun 0,1741 → 0,1414 setelah kalibrasi. "
         "Reliability curve memperlihatkan kedekatan prediksi dan frekuensi aktual; bukan jaminan tiap pelanggan.",
         .65, 6.1, 12, .8, size=17, color=MUTED)
    s = slide("Interpretasi fitur", "Koefisien model inti sebelum kalibrasi")
    picture(s, FIG / "coefficients.png", .6, 1.8, 8.0, 4.8)
    bullets(s, ["Tenure negatif: hubungan bersyarat dengan log-odds churn.", "Kontrak bulanan dan fiber punya bobot positif pada model inti.",
                "Koefisien MonthlyCharges negatif tidak otomatis salah.", "Fitur saling berhubungan; bobot bukan bukti kausal atau peringkat universal."],
            x=9, y=2, w=3.7, h=4.6, size=17)
    s = slide("Arsitektur & implementasi aplikasi", "Python • scikit-learn • pandas • joblib • Streamlit")
    for i, (title, content) in enumerate([
        ("INPUT", "Form 18 fitur\nCSV UTF-8\nTemplate unduhan"),
        ("VALIDASI", "Kolom & kategori\nAngka & rentang\nKonsistensi layanan"),
        ("MODEL", "Pipeline preprocess\nSigmoid calibrated\nProbabilitas churn"),
        ("OUTPUT", "Threshold 0,32\nCHURN/TETAP\nRisiko + CSV hasil"),
    ]):
        x = .65 + i*3.15
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.0), Inches(2.8), Inches(2.7))
        shape.fill.solid(); shape.fill.fore_color.rgb=rgb(CARD); shape.line.color.rgb=rgb(ACCENT)
        text(s, title, x+.15, 2.25, 2.5, .4, size=21, bold=True, color="A78BFA")
        text(s, content, x+.15, 3, 2.5, 1.5, size=20)
    text(s, "Artefak dan metadata diperiksa SHA-256. Hasil lama disembunyikan saat input berubah. "
         "CSV invalid ditolak dengan pesan dan nomor baris; tidak dibuang diam-diam.", .65, 5.5, 12, 1.1, size=20)
    s = slide("Demo aplikasi — jalankan langsung", "Buka URL live; tunjukkan interaksi, bukan hanya screenshot",
              "Gunakan template_pelanggan.csv. Klik Prediksi untuk contoh pelanggan, upload batch, unduh hasil, "
              "lalu upload CSV salah untuk menunjukkan validasi. Lihat docs/NASKAH_VIDEO_DEMO.md.")
    bullets(s, ["1. Buka aplikasi: KPI dan tab single, batch, dashboard, serta EDA.", "2. Isi profil pelanggan; klik Prediksi Sekarang dan jelaskan probabilitas/threshold.",
                "3. Download template → upload CSV → tampilkan ringkasan → download hasil.",
                "4. Tunjukkan validasi: CSV tanpa tenure atau kategori Contract yang salah.",
                "5. Buka dashboard: perbandingan model inti dan evaluasi model aplikasi."], size=21)
    text(s, APP_URL, .65, 6.35, 12, .4, size=13, color="60A5FA")
    s = slide("Kesimpulan", "Model terukur, aplikasi berjalan, hasil bisa direproduksi")
    bullets(s, ["LR dipilih berdasarkan CV-F1 training tertinggi (0,6304).",
                "Konfigurasi aplikasi: sigmoid calibration + threshold 0,32 train-only.",
                "Test aplikasi: accuracy 75,69%; precision 53,09%; recall 73,53%; F1 0,6166.",
                "275 dari 374 churn terdeteksi; 99 terlewat dan 243 false alarm.",
                "Nilai tambah: probabilitas konsisten, validasi input, batch CSV, dan dashboard evaluasi."], size=22)
    s = slide("Keterbatasan & pengembangan", "Kesimpulan disesuaikan dengan bukti yang tersedia")
    table(s, ["Keterbatasan", "Pengembangan"], [
        ["Dataset historis tunggal; pelanggan tenure 0 dieksklusi", "Validasi eksternal/temporal dan uji cohort pelanggan baru"],
        ["Profil identik lintas split; fitur TotalCharges dibuang tanpa ablation", "Group-aware split; ablation fitur; laporkan perubahan performa"],
        ["RF overfit; FP/FN masih cukup besar", "Regularisasi/tuning; alternatif boosting; threshold berbasis biaya nyata"],
        ["Retensi berbasis aturan; belum diuji efektivitasnya", "SHAP dan eksperimen tindakan retensi; monitoring drift"],
    ], h=4.2, size=19, widths=[.5,.5])
    s = slide("Pembagian tugas", "Diisi oleh anggota sesuai pengerjaan aktual — tanpa asumsi peran")
    table(s, ["Anggota", "NIM", "Tugas aktual"], [
        ["Muhamad Akhdan Ramadhan", "J0404241102", "[Isi sendiri]"],
        ["Thevan Erlangga", "J0404241073", "[Isi sendiri]"],
        ["Fachri Abyasa Tarid", "J0404241136", "[Isi sendiri]"],
    ], h=3.3, size=20, widths=[.43,.25,.32])
    text(s, "Seluruh anggota wajib memahami konsep, dataset, preprocessing, algoritma, evaluasi, dan aplikasi.",
         .65, 5.7, 12, .8, size=20, color=MUTED)
    s = slide("Referensi jurnal", "DOI dan teks lengkap tersedia di docs/STUDI_LITERATUR.md")
    for i, r in enumerate(refs):
        text(s, f"[{r['no']}] {r['short']} — {r['journal']}\n{r['doi']}",
             .7, 1.8+i*.9, 12, .8, size=18)
    s = slide("Source code & berkas pengumpulan", "Format folder: ML2026_B2_Kelompok12_PrediksiCustomerChurn")
    bullets(s, ["PPT: docs/" + PPT_NAME, "EDA: docs/eda/EDA.md • notebook: projekk.ipynb",
                "Lima jurnal & tabel: docs/STUDI_LITERATUR.md • docs/jurnal/",
                "Naskah demo ≤7 menit: docs/NASKAH_VIDEO_DEMO.md",
                "Pembagian tugas: isi pada slide 27; video asli direkam anggota."], size=21)
    text(s, "Repository: https://github.com/makhdanramadhan-ui/Machine-Learning-Project\n" + APP_URL,
         .65, 6.15, 12, .7, size=13, color="60A5FA")

    prs.save(DOCS / PPT_NAME)
    readable = []
    for s, item in zip(prs.slides, outlines):
        values = []
        for shape in s.shapes:
            if shape.has_text_frame:
                values.append(shape.text)
            elif shape.has_table:
                values.extend(" | ".join(c.text for c in row.cells) for row in shape.table.rows)
        readable.append(f"## Slide {item['slide']}: {item['title']}\n\n" + "\n\n".join(values) +
                        ("\n\nCatatan presenter: " + item['notes'] if item['notes'] else ""))
    (DOCS / "PPT_OUTLINE.md").write_text("# Outline PPT\n\n" + "\n\n".join(readable), encoding="utf-8")
    print(f"PPT: {len(prs.slides)} slide -> {DOCS / PPT_NAME}")
    print("Contoh demo:", example[["Prob_Churn","Prediksi"]].to_dict("records"))


if __name__ == "__main__":
    build_report()
